from __future__ import annotations

import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


NATIVE_FORMAT = "facebook-activity-native-json"
HEADERS = [
    "Type",
    "Tên người dùng",
    "Nội dung/Reaction",
    "Ngày",
    "Tiêu đề/Post",
    "Nơi bình luận",
    "Comment URL / Post URL",
    "IsComment",
    "IsReaction",
]
MEDIA_PATTERN = re.compile(r"\[(Hình ảnh|Video|Media):\s*([^\]]+)\]", re.IGNORECASE)
DIRECT_COMMENT_PATTERN = re.compile(
    r"comment_id|reply_comment_id|/comments?/|/comment/", re.IGNORECASE
)


def text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def number(value: Any) -> int:
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value)
    try:
        return int(text(value))
    except ValueError:
        return 0


def timestamp(value: Any) -> int:
    if isinstance(value, datetime):
        dt = value
    else:
        raw = text(value)
        if not raw:
            return 0
        dt = None
        for fmt in ("%H:%M:%S %d/%m/%Y", "%d/%m/%Y %H:%M:%S"):
            try:
                dt = datetime.strptime(raw, fmt)
                break
            except ValueError:
                continue
        if dt is None:
            return 0
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=ZoneInfo("Asia/Ho_Chi_Minh"))
    return int(dt.timestamp())


def media_from_content(value: str) -> tuple[str, list[dict[str, Any]]]:
    media: list[dict[str, Any]] = []
    for match in MEDIA_PATTERN.finditer(value):
        label = match.group(1).casefold()
        media_type = "IMAGE" if "hình" in label else "VIDEO" if "video" in label else "MEDIA"
        media.append({"type": media_type, "uri": text(match.group(2))})
    clean_content = text(MEDIA_PATTERN.sub(" ", value))
    return clean_content, media


def split_comment_url(url: str) -> tuple[str, str]:
    if not url:
        return "", ""
    if DIRECT_COMMENT_PATTERN.search(url):
        return url, ""
    return "", url


def source_row(sheet: str, row_number: int, values: list[Any], hyperlink: str) -> dict[str, Any]:
    cells = {header: values[index] if index < len(values) else None for index, header in enumerate(HEADERS)}
    return {
        "sheet": sheet,
        "rowNumber": row_number,
        "cells": cells,
        "urlHyperlink": hyperlink,
    }


def extract_fbid(url: str) -> str:
    match = re.search(r"/(?:permalink|posts?|comments?)/(\d+)", url, re.IGNORECASE)
    if match:
        return match.group(1)
    query_match = re.search(r"[?&](?:comment_id|reply_comment_id)=([^&#]+)", url, re.IGNORECASE)
    return query_match.group(1) if query_match else ""


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 3:
        raise SystemExit("Usage: convert_excel_to_native_json.py INPUT.xlsx OUTPUT.json")

    input_path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[2]).resolve()
    workbook = load_workbook(input_path, read_only=False, data_only=False)

    accounts: list[dict[str, Any]] = []
    skipped_rows: list[dict[str, Any]] = []
    totals = {
        "rows": 0,
        "comments": 0,
        "mediaComments": 0,
        "reactions": 0,
        "commentReactions": 0,
        "rowsWithUrl": 0,
        "rowsWithoutUrl": 0,
        "invalidDates": 0,
        "skippedRows": 0,
    }

    for sheet in workbook.worksheets:
        header_values = [text(sheet.cell(1, column).value) for column in range(1, 10)]
        account = {
            "accountName": sheet.title,
            "sourceSheet": sheet.title,
            "headers": header_values,
            "comments": [],
            "reactions": [],
        }

        for row_number in range(2, sheet.max_row + 1):
            cells = [sheet.cell(row_number, column) for column in range(1, 10)]
            values = [cell.value for cell in cells]
            row_type = text(values[0])
            normalized_type = row_type.casefold()
            if normalized_type not in {
                "comment",
                "comment - media",
                "reaction",
                "reaction comment",
            }:
                if any(value is not None for value in values):
                    skipped_rows.append(
                        source_row(sheet.title, row_number, values, text(cells[6].hyperlink.target if cells[6].hyperlink else ""))
                    )
                continue

            totals["rows"] += 1
            user_name = text(values[1]) or sheet.title
            content_or_reaction = text(values[2])
            date_text = text(values[3])
            activity_timestamp = timestamp(values[3])
            title = text(values[4])
            location = text(values[5])
            hyperlink = text(cells[6].hyperlink.target if cells[6].hyperlink else "")
            url = text(values[6]) or hyperlink
            raw_source = source_row(sheet.title, row_number, values, hyperlink)
            if activity_timestamp == 0 and date_text:
                totals["invalidDates"] += 1
            totals["rowsWithUrl" if url else "rowsWithoutUrl"] += 1

            if normalized_type.startswith("comment"):
                clean_content, media = media_from_content(content_or_reaction)
                comment_link, post_url = split_comment_url(url)
                account["comments"].append(
                    {
                        "authorName": user_name,
                        "content": clean_content,
                        "commentTime": activity_timestamp,
                        "commentId": extract_fbid(comment_link),
                        "title": title,
                        "group": location,
                        "commentLocation": location,
                        "commentLink": comment_link,
                        "commentSearchUrl": "",
                        "media": media,
                        "postUrl": post_url,
                        "source": raw_source,
                    }
                )
                totals["comments"] += 1
                if media or normalized_type == "comment - media":
                    totals["mediaComments"] += 1
            else:
                is_comment_reaction = normalized_type == "reaction comment"
                account["reactions"].append(
                    {
                        "reaction": content_or_reaction,
                        "linkPost": url,
                        "commentAuthorName": title or "Chưa xác định",
                        "ownerName": title or "Chưa xác định",
                        "reactionTime": activity_timestamp,
                        "fbid": extract_fbid(url),
                        "matchNames": [name for name in (location, title) if name],
                        "targetGroupNames": [location] if location else [],
                        "targetAuthorNames": [title] if title else [],
                        "targetType": "COMMENT" if is_comment_reaction else "POST",
                        "accountName": user_name,
                        "source": raw_source,
                    }
                )
                totals["reactions"] += 1
                if is_comment_reaction:
                    totals["commentReactions"] += 1

        if account["comments"] or account["reactions"]:
            accounts.append(account)

    totals["skippedRows"] = len(skipped_rows)
    result = {
        "format": NATIVE_FORMAT,
        "schemaVersion": 1,
        "createdAt": datetime.now(timezone.utc).isoformat(),
        "sourceWorkbook": {
            "fileName": input_path.name,
            "fileSize": input_path.stat().st_size,
            "modifiedAt": datetime.fromtimestamp(
                input_path.stat().st_mtime, tz=timezone.utc
            ).isoformat(),
            "sheetCount": len(workbook.sheetnames),
            "sheetNames": workbook.sheetnames,
            "expectedHeaders": HEADERS,
        },
        "summary": totals,
        "accounts": accounts,
        "skippedRows": skipped_rows,
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    print(json.dumps(totals, ensure_ascii=False))
    print(str(output_path))


if __name__ == "__main__":
    main()
